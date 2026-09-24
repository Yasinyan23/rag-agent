"""Live verification of all format-specific document loaders.

Generates minimal valid sample files, runs each through
``DocumentLoaderFactory.get_loader()`` and ``load()``, and prints a summary table.

Usage:
    python scripts/verify_all_loaders.py
"""

from __future__ import annotations

import asyncio
import csv
import io
import sys
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, StreamObject

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from src.ingestion.loaders.base import LoadedDocument  # noqa: E402
from src.ingestion.loaders.factory import DocumentLoaderFactory  # noqa: E402

_MARKER_TXT: str = "Plain text verification body."
_MARKER_MD_HEADING: str = "# Verification Markdown"
_MARKER_PDF: str = "Verification PDF sample text"
_MARKER_DOCX_HEADING: str = "# Sample DOCX Title"
_MARKER_CSV_PREFIX: str = "## CSV:"
_MARKER_SHEET_PREFIX: str = "## Sheet:"
_MARKER_HTML_HEADING: str = "# Loader Verification Page"


@dataclass(frozen=True)
class FormatProbe:
    """One generated sample and the section marker expected in loaded content."""

    label: str
    filename: str
    section_marker: str | None
    build_bytes: Callable[[], bytes]


def _build_sample_txt_bytes() -> bytes:
    return f"{_MARKER_TXT}\n".encode()


def _build_sample_md_bytes() -> bytes:
    return f"{_MARKER_MD_HEADING}\n\nMarkdown body for loader verification.\n".encode()


def _build_sample_pdf_bytes() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    content_stream = f"BT /F1 12 Tf 72 720 Td ({_MARKER_PDF}) Tj ET\n"
    stream = StreamObject()
    stream._data = content_stream.encode("latin-1")
    stream.update({NameObject("/Length"): NumberObject(len(stream._data))})
    page[NameObject("/Contents")] = stream
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _build_sample_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Sample DOCX Title", level=1)
    document.add_paragraph("DOCX body paragraph for verification.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "ColA"
    table.cell(0, 1).text = "ColB"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_sample_csv_bytes() -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Name", "Score"])
    writer.writerow(["Alpha", "10"])
    writer.writerow(["Beta", "20"])
    return buffer.getvalue().encode("utf-8")


def _build_sample_xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Metrics"
    sheet.append(["Metric", "Value"])
    sheet.append(["Latency", 42])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _build_sample_html_bytes() -> bytes:
    html = """<!DOCTYPE html>
<html lang="en">
  <head><title>Ignored When H1 Exists</title></head>
  <body>
    <h1>Loader Verification Page</h1>
    <p>HTML paragraph for loader verification.</p>
    <table>
      <tr><th>Key</th><th>Value</th></tr>
      <tr><td>Status</td><td>OK</td></tr>
    </table>
  </body>
</html>
"""
    return html.encode()


_PROBES: tuple[FormatProbe, ...] = (
    FormatProbe("Plain text (.txt)", "sample.txt", None, _build_sample_txt_bytes),
    FormatProbe("Markdown (.md)", "sample.md", _MARKER_MD_HEADING, _build_sample_md_bytes),
    FormatProbe("PDF (.pdf)", "sample.pdf", "## Page", _build_sample_pdf_bytes),
    FormatProbe("Word (.docx)", "sample.docx", _MARKER_DOCX_HEADING, _build_sample_docx_bytes),
    FormatProbe("CSV (.csv)", "sample.csv", _MARKER_CSV_PREFIX, _build_sample_csv_bytes),
    FormatProbe("Excel (.xlsx)", "sample.xlsx", _MARKER_SHEET_PREFIX, _build_sample_xlsx_bytes),
    FormatProbe("HTML (.html)", "sample.html", _MARKER_HTML_HEADING, _build_sample_html_bytes),
)


def _assert_loaded(probe: FormatProbe, loaded: LoadedDocument) -> str | None:
    """Return an error message when assertions fail, otherwise ``None``."""
    if not loaded.content.strip():
        return "LoadedDocument.content is empty"

    if probe.section_marker is not None and probe.section_marker not in loaded.content:
        return f"Expected section marker {probe.section_marker!r} not found in content"

    if probe.filename == "sample.txt" and _MARKER_TXT not in loaded.content:
        return f"Expected body marker {_MARKER_TXT!r} not found in content"

    return None


async def _run_probe(probe: FormatProbe, file_bytes: bytes) -> tuple[bool, str]:
    extension = Path(probe.filename).suffix
    loader = DocumentLoaderFactory.get_loader(extension)
    try:
        loaded = await loader.load(file_bytes, probe.filename)
    except Exception as exc:
        return False, str(exc)

    error = _assert_loaded(probe, loaded)
    if error is not None:
        return False, error
    return True, f"{len(loaded.content)} chars"


def _print_summary_table(rows: list[tuple[str, str, str, str]]) -> None:
    headers = ("Format", "File", "Status", "Detail")
    widths = [len(h) for h in headers]
    for row in rows:
        for index, cell in enumerate(row):
            widths[index] = max(widths[index], len(cell))

    def _fmt_row(cells: tuple[str, ...]) -> str:
        return "  ".join(cell.ljust(widths[i]) for i, cell in enumerate(cells))

    separator = "  ".join("-" * widths[i] for i in range(len(headers)))
    print()
    print(_fmt_row(headers))
    print(separator)
    for row in rows:
        print(_fmt_row(row))
    print()


async def _verify_all() -> int:
    table_rows: list[tuple[str, str, str, str]] = []
    failures = 0

    with TemporaryDirectory(prefix="docuquery_loader_verify_") as temp_dir:
        temp_path = Path(temp_dir)
        for probe in _PROBES:
            file_bytes = probe.build_bytes()
            disk_path = temp_path / probe.filename
            disk_path.write_bytes(file_bytes)
            read_back = disk_path.read_bytes()
            ok, detail = await _run_probe(probe, read_back)
            status = "PASS" if ok else "FAIL"
            if not ok:
                failures += 1
            table_rows.append((probe.label, probe.filename, status, detail))

    _print_summary_table(table_rows)

    if failures:
        print(f"{failures} format(s) failed verification.", file=sys.stderr)
        return 1

    print("All document loader formats verified successfully.")
    return 0


def main() -> None:
    exit_code = asyncio.run(_verify_all())
    sys.exit(exit_code)


if __name__ == "__main__":
    main()
