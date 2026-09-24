"""Integration tests for document uploads on ``POST /api/v1/ingest/file``."""

from __future__ import annotations

import io
from unittest.mock import AsyncMock

import pytest
from docx import Document
from httpx import AsyncClient
from openpyxl import Workbook


def _minimal_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Report", level=1)
    doc.add_paragraph("Sample DOCX body for ingestion.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _minimal_html_bytes() -> bytes:
    return b"""<!DOCTYPE html>
<html><head><title>Report</title></head>
<body><h1>Report</h1><p>Sample HTML body for ingestion.</p></body></html>"""


def _minimal_xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["metric", "value"])
    sheet.append(["latency_ms", 12])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_ingest_processes_pdf_file(
    integration_client: AsyncClient,
    mock_ingestion_pipeline: AsyncMock,
) -> None:
    """``POST /api/v1/ingest/file`` must accept ``.pdf`` files and return HTTP 200."""
    mock_ingestion_pipeline.ingest_document.return_value = 5

    response = await integration_client.post(
        "/api/v1/ingest/file",
        files={"file": ("manual.pdf", b"%PDF-1.4 sample", "application/pdf")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert payload["filename"] == "manual.pdf"
    assert payload["chunks_ingested"] == 5


@pytest.mark.asyncio
async def test_ingest_processes_docx_file(
    integration_client: AsyncClient,
    mock_ingestion_pipeline: AsyncMock,
) -> None:
    """``POST /api/v1/ingest/file`` must accept ``.docx`` files and return HTTP 200."""
    mock_ingestion_pipeline.ingest_document.return_value = 4

    response = await integration_client.post(
        "/api/v1/ingest/file",
        files={
            "file": (
                "report.docx",
                _minimal_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert payload["filename"] == "report.docx"
    assert payload["chunks_ingested"] == 4


@pytest.mark.asyncio
async def test_ingest_processes_html_file(
    integration_client: AsyncClient,
    mock_ingestion_pipeline: AsyncMock,
) -> None:
    """``POST /api/v1/ingest/file`` must accept ``.html`` files and return HTTP 200."""
    mock_ingestion_pipeline.ingest_document.return_value = 2

    response = await integration_client.post(
        "/api/v1/ingest/file",
        files={"file": ("guide.html", _minimal_html_bytes(), "text/html")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert payload["filename"] == "guide.html"
    assert payload["chunks_ingested"] == 2


@pytest.mark.asyncio
async def test_ingest_processes_csv_file(
    integration_client: AsyncClient,
    mock_ingestion_pipeline: AsyncMock,
) -> None:
    """``POST /api/v1/ingest/file`` must accept ``.csv`` files and return HTTP 200."""
    mock_ingestion_pipeline.ingest_document.return_value = 3

    response = await integration_client.post(
        "/api/v1/ingest/file",
        files={"file": ("metrics.csv", b"id,value\n1,42\n", "text/csv")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert payload["filename"] == "metrics.csv"
    assert payload["chunks_ingested"] == 3


@pytest.mark.asyncio
async def test_ingest_processes_xlsx_file(
    integration_client: AsyncClient,
    mock_ingestion_pipeline: AsyncMock,
) -> None:
    """``POST /api/v1/ingest/file`` must accept ``.xlsx`` files and return HTTP 200."""
    mock_ingestion_pipeline.ingest_document.return_value = 6

    response = await integration_client.post(
        "/api/v1/ingest/file",
        files={
            "file": (
                "workbook.xlsx",
                _minimal_xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert payload["filename"] == "workbook.xlsx"
    assert payload["chunks_ingested"] == 6
