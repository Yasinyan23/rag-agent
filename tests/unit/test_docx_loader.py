"""Unit tests for the DOCX document loader."""

from __future__ import annotations

import io

import pytest
from docx import Document

from src.core.exceptions import DocumentParsingError
from src.ingestion.chunker import TokenSlidingWindowChunker
from src.ingestion.loaders import DocxDocumentLoader as DocxLoaderExport
from src.ingestion.loaders.docx import DocxDocumentLoader
from src.ingestion.loaders.factory import DocumentLoaderFactory


def _build_sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Heading 1", level=1)
    doc.add_paragraph("Body under the first heading.")
    doc.add_heading("Heading 2", level=2)
    doc.add_paragraph("Body under the second heading.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Alpha"
    table.cell(0, 1).text = "Beta"
    table.cell(1, 0).text = "Gamma"
    table.cell(1, 1).text = "Delta"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_empty_docx_bytes() -> bytes:
    doc = Document()
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_public_exports_include_docx_loader() -> None:
    assert DocxLoaderExport is DocxDocumentLoader


@pytest.mark.asyncio
async def test_docx_loader_extracts_headings_paragraphs_and_tables() -> None:
    loaded = await DocxDocumentLoader().load(_build_sample_docx_bytes(), "sample.docx")

    assert "# Heading 1" in loaded.content
    assert "## Heading 2" in loaded.content
    assert "Body under the first heading." in loaded.content
    assert "Alpha | Beta" in loaded.content
    assert "Gamma | Delta" in loaded.content
    assert loaded.source == "sample"
    assert loaded.metadata["format"] == "docx"
    assert loaded.metadata["table_count"] == 1


@pytest.mark.asyncio
async def test_docx_loader_rejects_empty_file_bytes() -> None:
    with pytest.raises(DocumentParsingError, match="empty"):
        await DocxDocumentLoader().load(b"   ", "empty.docx")


@pytest.mark.asyncio
async def test_docx_loader_rejects_corrupted_docx() -> None:
    with pytest.raises(DocumentParsingError, match="corrupted|valid"):
        await DocxDocumentLoader().load(b"not-a-valid-docx", "broken.docx")


@pytest.mark.asyncio
async def test_docx_loader_rejects_docx_with_no_extractable_text() -> None:
    with pytest.raises(DocumentParsingError, match="no extractable content"):
        await DocxDocumentLoader().load(_build_empty_docx_bytes(), "blank.docx")


def _build_sectioned_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Heading 1", level=1)
    doc.add_paragraph("Body under the first heading. " * 50)
    doc.add_heading("Heading 2", level=2)
    doc.add_paragraph("Body under the second heading. " * 50)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_docx_heading_maps_to_chunk_section_metadata() -> None:
    loaded = await DocxDocumentLoader().load(_build_sectioned_docx_bytes(), "sections.docx")

    chunker = TokenSlidingWindowChunker(chunk_size=40, chunk_overlap=5)
    chunks = chunker.chunk_document("sections", loaded.content)
    sections = {chunk.metadata["section"] for chunk in chunks if chunk.metadata["section"]}

    assert "Heading 1" in sections
    assert "Heading 2" in sections


def test_factory_resolves_docx_from_filename() -> None:
    assert isinstance(DocumentLoaderFactory.get_loader("sample.docx"), DocxDocumentLoader)
    assert isinstance(DocumentLoaderFactory.get_loader(".docx"), DocxDocumentLoader)
