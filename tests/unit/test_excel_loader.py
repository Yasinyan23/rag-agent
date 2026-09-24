"""Unit tests for the XLSX document loader."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from src.core.exceptions import DocumentParsingError
from src.ingestion.chunker import TokenSlidingWindowChunker
from src.ingestion.loaders import ExcelDocumentLoader as ExcelLoaderExport
from src.ingestion.loaders.excel import ExcelDocumentLoader
from src.ingestion.loaders.factory import DocumentLoaderFactory


def test_public_exports_include_excel_loader() -> None:
    assert ExcelLoaderExport is ExcelDocumentLoader


def _build_multi_sheet_workbook_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Sales"
    first.append(["Region", "Revenue"])
    first.append(["North", 100])

    second = workbook.create_sheet("Costs")
    second.append(["Item", "Amount"])
    second.append(["Hosting", 50])

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _build_empty_workbook_bytes() -> bytes:
    workbook = Workbook()
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_excel_loader_extracts_visible_sheets_as_markdown() -> None:
    loaded = await ExcelDocumentLoader().load(_build_multi_sheet_workbook_bytes(), "report.xlsx")

    assert "## Sheet: Sales" in loaded.content
    assert "## Sheet: Costs" in loaded.content
    assert "| Region | Revenue |" in loaded.content
    assert "| North | 100 |" in loaded.content
    assert "| Item | Amount |" in loaded.content
    assert loaded.source == "report"
    assert loaded.metadata["format"] == "xlsx"
    assert loaded.metadata["sheet_count"] == 2


@pytest.mark.asyncio
async def test_excel_loader_rejects_empty_workbook() -> None:
    with pytest.raises(DocumentParsingError, match="no extractable rows"):
        await ExcelDocumentLoader().load(_build_empty_workbook_bytes(), "empty.xlsx")


@pytest.mark.asyncio
async def test_excel_loader_rejects_corrupt_file() -> None:
    with pytest.raises(DocumentParsingError, match="corrupted|valid"):
        await ExcelDocumentLoader().load(b"not-a-valid-xlsx", "broken.xlsx")


@pytest.mark.asyncio
async def test_excel_loader_rejects_empty_file_bytes() -> None:
    with pytest.raises(DocumentParsingError, match="empty"):
        await ExcelDocumentLoader().load(b"   ", "empty.xlsx")


def _build_sectioned_workbook_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Sales"
    first.append(["Region", "Notes"])
    first.append(["North", " ".join(["revenue"] * 80)])

    second = workbook.create_sheet("Costs")
    second.append(["Item", "Notes"])
    second.append(["Hosting", " ".join(["spend"] * 80)])

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_excel_sheet_heading_maps_to_chunk_section_metadata() -> None:
    loaded = await ExcelDocumentLoader().load(_build_sectioned_workbook_bytes(), "report.xlsx")

    chunker = TokenSlidingWindowChunker(chunk_size=40, chunk_overlap=5)
    chunks = chunker.chunk_document("report", loaded.content)
    sections = {chunk.metadata["section"] for chunk in chunks if chunk.metadata["section"]}

    assert "Sheet: Sales" in sections
    assert "Sheet: Costs" in sections


def test_factory_resolves_xlsx_from_filename() -> None:
    assert isinstance(DocumentLoaderFactory.get_loader("report.xlsx"), ExcelDocumentLoader)
    assert isinstance(DocumentLoaderFactory.get_loader(".xlsx"), ExcelDocumentLoader)
