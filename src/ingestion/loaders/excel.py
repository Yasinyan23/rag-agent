"""XLSX document loader with per-sheet Markdown sections for citation provenance."""

from __future__ import annotations

import asyncio
import io
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.core.exceptions import DocumentParsingError
from src.ingestion.loaders.base import BaseDocumentLoader, LoadedDocument
from src.ingestion.loaders.csv_loader import _normalize_row, _rows_to_markdown_table


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _sheet_rows(sheet: Worksheet) -> list[list[str]]:
    rows: list[list[str]] = []
    for raw_row in sheet.iter_rows(values_only=True):
        if raw_row is None:
            continue
        cells = [_cell_to_str(cell) for cell in raw_row]
        if not any(cell.strip() for cell in cells):
            continue
        rows.append(_normalize_row(cells))
    return rows


def _extract_excel_text(file_bytes: bytes, filename: str) -> LoadedDocument:
    """Synchronous XLSX parse — intended to run inside ``asyncio.to_thread``."""
    source = Path(filename).stem
    try:
        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except BadZipFile as exc:
        raise DocumentParsingError(
            f"Unable to read XLSX {filename!r}: the file may be corrupted or not a valid "
            f"Excel workbook."
        ) from exc
    except (KeyError, ValueError, OSError) as exc:
        raise DocumentParsingError(
            f"Unable to read XLSX {filename!r}: the file may be corrupted or invalid."
        ) from exc

    blocks: list[str] = []
    sheet_count = 0
    total_row_count = 0

    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue

            rows = _sheet_rows(sheet)
            if not rows:
                continue

            table_block = _rows_to_markdown_table(rows)
            if not table_block.strip():
                continue

            sheet_count += 1
            total_row_count += len(rows)
            sheet_title = sheet.title or f"Sheet{sheet_count}"
            blocks.append(f"## Sheet: {sheet_title}\n\n{table_block}")
    finally:
        workbook.close()

    if not blocks:
        raise DocumentParsingError(
            f"Unable to extract data from XLSX {filename!r}: workbook contains no extractable rows."
        )

    metadata: dict[str, Any] = {
        "format": "xlsx",
        "filename": filename,
        "sheet_count": sheet_count,
        "row_count": total_row_count,
    }

    return LoadedDocument(
        content="\n\n".join(blocks),
        source=source,
        metadata=metadata,
    )


class ExcelDocumentLoader(BaseDocumentLoader):
    """Extracts visible XLSX sheets as Markdown tables with sheet-level headings."""

    async def load(self, file_bytes: bytes, filename: str) -> LoadedDocument:
        """Parse ``file_bytes`` off the event loop and return sheet-marked tabular text."""
        if not file_bytes.strip():
            raise DocumentParsingError(f"Unable to read XLSX {filename!r}: file is empty.")

        return await asyncio.to_thread(_extract_excel_text, file_bytes, filename)
